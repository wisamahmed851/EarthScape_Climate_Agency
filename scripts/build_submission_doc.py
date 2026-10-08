"""Insert the redesigned dashboard screenshots and as-built notes into a COPY of the team's project document"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "docs" / "screenshots"
BODY_FONT = "Book Antiqua"


def main(src):
    doc = Document(src)
    anchor = next(p for p in doc.paragraphs if p.style.name == "Heading 1" and p.text.strip() == "Testing & Troubleshooting")

    def para(text="", style=None, size=None, bold=False, italic=False, font=BODY_FONT, color=None, align=None, space_after=None):
        p = anchor.insert_paragraph_before(style=style)
        if text:
            r = p.add_run(text)
            r.font.name, r.bold, r.italic = font, bold, italic
            r._element.rPr.rFonts.set(qn("w:hAnsi"), font)
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
        if align:
            p.alignment = align
        if space_after is not None:
            p.paragraph_format.space_after = Pt(space_after)
        return p

    def h1(text):
        p = para(text, style="Heading 1", size=24, font="Berlin Sans FB Demi", color="000000")
        p.paragraph_format.page_break_before = True
        return p

    def h2(text):
        return para(text, style="Heading 2", size=16, bold=True, color="2C807C")

    def shot(name, caption, width=17.5):
        p = para(align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
        p.add_run().add_picture(str(SHOTS / name), width=Cm(width))
        p.paragraph_format.keep_with_next = True
        para(caption, size=11, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

    def table(header, rows, widths):
        t = doc.add_table(rows=1, cols=len(header))
        t.style = doc.tables[0].style
        for i, text in enumerate(header):
            t.rows[0].cells[i].text = text
        for r in rows:
            cells = t.add_row().cells
            for i, text in enumerate(r):
                cells[i].text = text
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
            for c in row.cells:
                for p in c.paragraphs:
                    for run in p.runs:
                        run.font.size, run.font.name = Pt(10.5), "Open Sans"
        anchor._p.addprevious(t._tbl)
        para(space_after=6)

    h1("Dashboard Implementation")
    para("The EarthScape dashboard is a web application (FastAPI, MongoDB, Chart.js) that follows the project identity: the EarthScape logo in the top bar and on the sign-in page, "
         "the deep-teal and navy palette of this document, and the same serif headings. It is reached at http://127.0.0.1:8000 and is bound to the local machine only. "
         "The screenshots below were taken from the running system with real stored data; labels on every page say whether a value is gridded reanalysis (NASA POWER), modelled (Open-Meteo) or observed (OpenAQ).", size=14)

    h2("Sign-in and role-based access")
    para("Users sign in with a personal account. Administrators additionally see User Management and Feedback Review; Analysts receive a permission-denied page for those routes.", size=13)
    shot("ui_login.png", "Figure A. Sign-in page with the EarthScape logo", 14)
    shot("ui_users.png", "Figure B. User Management (Administrator only): create users, change roles, disable accounts, reset passwords")

    h2("Historical climate (NASA POWER, processed with Hadoop MapReduce)")
    shot("ui_overview.png", "Figure C. Overview: key figures for the chosen city with the data-source label")
    shot("ui_history.png", "Figure D. Historical Climate: temperature, precipitation, humidity and wind with city, date and resolution filters")
    shot("ui_compare.png", "Figure E. City Comparison across Karachi, Lahore, Islamabad, Peshawar and Quetta")

    h2("Analytics: trends, anomalies, correlation and forecasts")
    shot("ui_trends.png", "Figure F. Trends and Extremes: Theil-Sen trends, anomaly counts and extreme-event counts from the yearly MapReduce job")
    shot("ui_correlation.png", "Figure G. Correlation of deseasonalised daily variables (association, not causation)")
    shot("ui_forecast.png", "Figure H. Forecast: model-derived values, shown apart from observations, with held-out test results")

    h2("Live data, historical comparison and alerts")
    shot("ui_current.png", "Figure I. Current Data: source freshness, modelled and observed values, and the current-versus-historical comparison (top of page)")
    shot("ui_alerts.png", "Figure J. Alert Rules: configurable thresholds on live variables")
    shot("ui_alerts_history.png", "Figure K. Alert History: raised alerts with value, unit, source and acknowledgement")

    h2("Provenance, feedback and mobile use")
    shot("ui_sources.png", "Figure L. Data Sources: provider classes, provenance and pipeline status")
    shot("ui_feedback.png", "Figure M. Feedback and support form")
    shot("ui_mobile_overview.png", "Figure N. The dashboard on a 390 px wide phone screen", 7.5)

    h1("As-Built Implementation Notes")
    para("The proposed designs earlier in this document asked for the collection names and algorithms to be aligned with the final implementation. The tables below describe what was actually built.", size=14)
    h2("Database design as built (MongoDB database earthscape)")
    table(["Collection", "Purpose", "Relation to the proposed design"], [
        ["users", "accounts: username, Argon2id password hash, role (Administrator / Analyst), active flag, lockout data", "USERS (no e-mail field)"],
        ["sessions", "server-side login sessions (token hash, CSRF token, expiry)", "added for authentication"],
        ["alert_rules", "city, variable, operator, threshold, severity, enabled", "ALERT RULES"],
        ["alerts", "raised alerts: value, unit, source, modelled/observed tag, status, acknowledged by", "ALERT EVENTS (no delivery_status; alerts are in-app only)"],
        ["latest_readings", "normalised live readings from Open-Meteo and OpenAQ, unique per source/city/station/time, 30-day expiry", "stands in for DATASETS of live sources"],
        ["ingest_status", "last attempt, success, failures and cycle duration per source", "monitoring (not in the proposal)"],
        ["ml_runs", "versioned analytics runs: trends, anomalies, correlation, forecasts, seed", "MODEL VERSIONS and PREDICTION RESULTS combined"],
        ["feedback", "support requests and feedback with status", "added for the feedback requirement"],
    ], [3.3, 9.0, 6.0])
    para("DATA SOURCES, DATASETS and PROCESSING JOBS for the historical data are not MongoDB collections: they are append-only provenance manifests in HDFS (/earthscape/_manifest/*.jsonl) holding source, size, SHA-256, record counts and the YARN job and application identifiers. "
         "Raw, interim and processed files are in HDFS and are not duplicated in MongoDB.", size=13)
    h2("Machine-learning methods as built")
    table(["Proposed", "Implemented", "Note"], [
        ["Linear Regression (trend)", "Theil-Sen slope with Mann-Kendall test for trends; Ridge regression on a linear trend plus annual harmonics for forecasts", "Ridge is regularised linear regression; chosen by rolling-origin validation against seasonal baselines"],
        ["Random Forest Regression", "not implemented", "no experiment was run; nothing is claimed for it"],
        ["Isolation Forest (anomaly)", "Isolation Forest plus a robust z-score rule, both fitted on 2001-2015 only", "no labelled anomalies exist, so precision and recall are not reported"],
        ["Correlation", "Pearson and Spearman on deseasonalised daily anomalies", "association only"],
    ], [4.2, 7.6, 6.5])
    para("Other differences from the proposal that the team should know: near-real-time REST polling replaces event streaming; alerts are shown in the application (no e-mail or SMS); MODIS satellite data and weather-station records are not ingested; "
         "TLS and encryption at rest are prepared or documented but not active; there is no dataset upload page (data enter through the ingestion scripts); and a dedicated audit trail of logins, role changes and configuration changes is not implemented.", size=13)

    settings = doc.settings.element
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    settings.append(uf)

    out = ROOT / "submission" / "EarthScape Climate Agency - with dashboard.docx"
    out.parent.mkdir(exist_ok=True)
    doc.save(out)
    print("written", out)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\AKL\Downloads\EarthScape Climate Agency.docx")
