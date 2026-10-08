r"""Build the SRS deliverables as genuine Word documents: submission/EarthScape_Final_Report.docx and submission/ReadMe.docx.

Run: .venv\Scripts\python.exe scripts\build_report.py [--tests N]
Facts come from the project documents (docs/*.md); the requirement table is read from docs/srs-completion-plan.md.
The SRS asks for "ReadMe.doc": this script writes .docx (the format available here); see the ReadMe for the Save-As step.
"""
import argparse
import re
from datetime import date
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "docs" / "screenshots"
OUT = ROOT / "submission"


def diagram(path):
    """Architecture / data-flow figure drawn with matplotlib (no external service)."""
    fig, ax = plt.subplots(figsize=(11, 6.2))
    ax.set_xlim(0, 22)
    ax.set_ylim(0, 12.4)
    ax.axis("off")
    colors = {"src": "#E8F1FA", "hdfs": "#FFF3D6", "proc": "#E6F4EA", "app": "#F3E8FF", "db": "#FDE8E8"}

    def box(x, y, w, h, text, kind):
        ax.add_patch(plt.Rectangle((x, y), w, h, fc=colors[kind], ec="#333", lw=1))
        ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=6.8)

    def arrow(x1, y1, x2, y2):
        ax.annotate("", xy=(x2, y2), xytext=(x1, y1), arrowprops=dict(arrowstyle="->", color="#333", lw=1))

    ax.text(0.2, 11.9, "Batch layer (historical)", fontsize=9, weight="bold")
    box(0.2, 9.8, 3.4, 1.6, "NASA POWER hourly\n2001-2025 (gridded\nreanalysis, 5 cities)", "src")
    box(4.6, 9.8, 3.2, 1.6, "HDFS RAW\n125 CSV, immutable\nmanifest + sha256", "hdfs")
    box(8.8, 9.8, 3.2, 1.6, "HDFS INTERIM\n5 CSV, 1.1 M rows\n-999 -> null + flag", "hdfs")
    box(13.0, 9.8, 4.0, 1.6, "Hadoop MapReduce on YARN\ndaily / monthly / yearly\n(Streaming, Python)", "proc")
    box(18.0, 9.8, 3.6, 1.6, "HDFS PROCESSED\n45,655 / 1,500 / 125 rows\n+ provenance", "hdfs")
    for x1, x2 in ((3.6, 4.6), (7.8, 8.8), (12.0, 13.0), (17.0, 18.0)):
        arrow(x1, 10.6, x2, 10.6)
    ax.text(0.2, 8.6, "Near-real-time layer (REST polling, not streaming)", fontsize=9, weight="bold")
    box(0.2, 6.5, 3.4, 1.6, "Open-Meteo weather +\nair quality\n(MODELLED)", "src")
    box(0.2, 4.5, 3.4, 1.6, "OpenAQ PM2.5\n(OBSERVED sensors)", "src")
    box(4.6, 5.5, 3.2, 1.8, "Poller thread\n15 min / OpenAQ 60 min\nretries, stale rules", "proc")
    box(8.8, 6.5, 3.2, 1.6, "MongoDB latest_readings\nunique per source/city/\nstation/time", "db")
    box(8.8, 4.3, 3.2, 1.6, "Daily sealed RAW\nHDFS date= partitions\nsha256 + manifest", "hdfs")
    box(13.0, 6.5, 4.0, 1.6, "Alert evaluator\nthresholds, stale-safe,\nno duplicate alerts", "proc")
    box(13.0, 4.3, 4.0, 1.6, "ML (trend, anomaly,\ncorrelation, forecast)\nversioned runs", "proc")
    for y in (7.3, 5.3):
        arrow(3.6, y, 4.6, 6.4)
    arrow(7.8, 6.4, 8.8, 7.3)
    arrow(7.8, 6.0, 8.8, 5.1)
    arrow(12.0, 7.3, 13.0, 7.3)
    arrow(19.8, 9.8, 15.0, 5.9)
    ax.text(0.2, 3.3, "Application and operations", fontsize=9, weight="bold")
    box(0.2, 0.8, 5.2, 2.0, "FastAPI dashboards\nlogin, Administrator / Analyst,\nCSRF, 127.0.0.1:8000", "app")
    box(6.4, 0.8, 4.2, 2.0, "MongoDB: users, sessions,\nfeedback, rules, alerts,\nml_runs, ingest_status", "db")
    box(11.6, 0.8, 4.6, 2.0, "Windows Task Scheduler:\nhealth check every 15 min,\ndaily backup (keep 7)", "proc")
    box(17.2, 0.8, 4.4, 2.0, "Current vs historical:\nPOWER hour climatology\nvs Open-Meteo (T, RH)", "app")
    for x in (15.0, 10.4):
        arrow(x, 4.3, 3.5, 2.8)
    fig.savefig(path, dpi=150, bbox_inches="tight")
    plt.close(fig)


class Doc:
    def __init__(self, title):
        self.d = Document()
        s = self.d.sections[0]
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2.0)
        st = self.d.styles["Normal"]
        st.font.name, st.font.size = "Calibri", Pt(10.5)
        for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
            h = self.d.styles[name]
            h.font.name, h.font.size, h.font.color.rgb = "Calibri", Pt(size), RGBColor(0x1F, 0x3A, 0x5F)
        footer = s.footer.paragraphs[0]
        footer.text = f"{title} - {date.today():%d %B %Y}"
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def h(self, text, level=1):
        self.d.add_heading(text, level)

    def p(self, text, bold=False, italic=False):
        para = self.d.add_paragraph()
        for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):   # **bold** inline
            run = para.add_run(part)
            run.bold = bold or i % 2 == 1
            run.italic = italic
        return para

    def bullets(self, items):
        for t in items:
            para = self.d.add_paragraph(style="List Bullet")
            for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", t)):
                para.add_run(part).bold = i % 2 == 1

    def numbered(self, items):
        for t in items:
            self.d.add_paragraph(t, style="List Number")

    def code(self, text):
        for line in text.strip("\n").split("\n"):
            para = self.d.add_paragraph()
            para.paragraph_format.space_after = Pt(0)
            run = para.add_run(line)
            run.font.name, run.font.size = "Consolas", Pt(8.5)

    def table(self, header, rows, widths=None, size=8.5):
        t = self.d.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, text in enumerate(header):
            c = t.rows[0].cells[i]
            c.text = ""
            run = c.paragraphs[0].add_run(str(text))
            run.bold, run.font.size = True, Pt(size)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:val"), "clear")
            shade.set(qn("w:fill"), "D9E2F3")
            c._tc.get_or_add_tcPr().append(shade)
        for r in rows:
            cells = t.add_row().cells
            for i, text in enumerate(r):
                cells[i].text = ""
                run = cells[i].paragraphs[0].add_run(re.sub(r"[`*]", "", str(text)))
                run.font.size = Pt(size)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        self.d.add_paragraph()

    def image(self, path, caption, width=16.0):
        if Path(path).exists():
            self.d.add_picture(str(path), width=Cm(width))
            self.d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap = self.d.add_paragraph(caption)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.runs[0].italic = True
            cap.runs[0].font.size = Pt(9)

    def pagebreak(self):
        self.d.add_page_break()

    def save(self, path):
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        self.d.save(path)


def requirement_rows():
    rows = []
    for line in (ROOT / "docs" / "srs-completion-plan.md").read_text(encoding="utf-8").splitlines():
        if re.match(r"\| [A-Z]+-\d+ \|", line):
            cells = [c.strip() for c in line.strip().strip("|").split(" | ")]
            if len(cells) == 6:
                cells[2] = cells[2].strip("*")
                rows.append(cells)
    return rows


def final_report(tests):
    OUT.mkdir(exist_ok=True)
    fig = SHOTS / "architecture_diagram.png"
    diagram(fig)
    d = Doc("EarthScape Climate Agency - Final Project Report")
    d.d.add_heading("EarthScape Climate Agency", 0)
    d.p("Big-data climate and air-quality monitoring for Karachi, Lahore, Islamabad, Peshawar and Quetta", bold=True)
    d.p("Final project report (Aptech eProject, Big Data). Prepared " + f"{date.today():%d %B %Y}" + ". Student name, roll number and centre: to be completed by the student.")
    d.p("**Honesty statement.** Everything reported as working was run and verified; items that could not be completed are listed in Section 14 and in the requirement table (Section 4). "
        "No data, metric or test result in this report is invented. The demonstration video has **not** been recorded by this document's author.")

    d.h("1. Problem definition and objectives")
    d.p("The EarthScape Climate Agency monitors climate change. It needs a solution that collects climate data from several kinds of source, stores it scalably, processes it in parallel, applies machine learning, "
        "visualises it interactively for authenticated users and alerts stakeholders when thresholds are crossed. The SRS requires HDFS storage and Hadoop MapReduce processing.")
    d.bullets([
        "Integrate approved sources: 25 years of NASA POWER hourly reanalysis, Open-Meteo current weather and air quality (modelled), OpenAQ PM2.5 (observed) and MODIS land-surface temperature (planned).",
        "Store the data in HDFS (RAW, INTERIM, PROCESSED) with partitioning, manifests and checksums.",
        "Process the historical data with Hadoop MapReduce on YARN (daily, monthly and yearly aggregates).",
        "Provide near-real-time monitoring by REST polling and threshold alerts.",
        "Provide trend analysis, anomaly detection, correlation analysis and forecasting that can be re-trained.",
        "Provide role-protected interactive dashboards, a feedback system, monitoring, backups and documentation.",
    ])

    d.h("2. Scope, assumptions and technologies")
    d.table(["Item", "Decision"], [
        ["Geographic scope", "Pakistan: Karachi, Lahore, Islamabad, Peshawar, Quetta (coordinates in config/cities.json)"],
        ["Historical data", "NASA POWER hourly, 2001-01-01 to 2025-12-31: T2M, RH2M, PRECTOTCORR, WS2M, PS, ALLSKY_SFC_SW_DWN (gridded reanalysis, not station data)"],
        ["Current data", "Open-Meteo weather and air quality (model output, REST polling); OpenAQ PM2.5 (observed by third-party monitors)"],
        ["Satellite data", "MODIS MOD11A2 v061 LST, tiles h24v06/h24v05/h23v05 (not ingested: no NASA Earthdata credentials)"],
        ["Big data stack", "Hadoop 3.4.3 (HDFS, YARN, MapReduce via Hadoop Streaming in Python) in WSL2 Ubuntu 24.04"],
        ["Application", "Python 3.13, FastAPI, Jinja2, Bootstrap, Chart.js, MongoDB 8.2, Argon2id, scikit-learn, SciPy, pandas"],
        ["Real-time", "Near-real-time REST polling thread (15 min; OpenAQ hourly). No Kafka or Spark."],
        ["Notifications", "In-app alerts only (no email or SMS)"],
        ["Deployment", "Local laptop, bound to 127.0.0.1; not public"],
    ], [4.0, 12.6])

    d.h("3. Requirement sources")
    d.p("The requirements come from the Aptech SRS 'Big Data - EarthScape Climate Agency' (Functional Requirements, Non-Functional Requirements, Hardware/Software Requirements, Project Deliverables). "
        "The SRS file itself is not redistributed in the submission package. Software named by the SRS without a stated function (Apache server, Tableau, Impala server, RStudio, Anaconda) was not treated as optional: each has its own line in Section 4.")

    d.h("4. Requirement traceability (40 items)")
    rows = requirement_rows()
    counts = {}
    for r in rows:
        key = r[2].split(" (")[0]
        counts[key] = counts.get(key, 0) + 1
    d.p("Status summary: " + ", ".join(f"{v} {k}" for k, v in sorted(counts.items(), key=lambda kv: -kv[1])) + f" (total {len(rows)}). Complete means the code or data exists and was run and verified.")
    d.table(["ID", "SRS wording / section", "Status", "Evidence", "Missing / limits", "User action"],
            rows, [1.2, 4.0, 1.5, 4.6, 3.2, 2.3], size=7)

    d.h("5. System architecture")
    d.p("Single-machine academic system: Windows 11 with WSL2 for Hadoop. There are no microservices, containers or queues. The figure shows the batch layer, the near-real-time layer and the application and operations components.")
    d.image(fig, "Figure 1. EarthScape architecture and data flow", 16.5)
    d.table(["Component", "Technology", "Role"], [
        ["Web app", "FastAPI + Jinja2, Bootstrap, Chart.js, uvicorn on 127.0.0.1:8000", "dashboards, JSON APIs, authentication, administration"],
        ["Application database", "MongoDB (database earthscape)", "users, sessions, feedback, alert rules and alerts, current readings, ingest status, ML runs"],
        ["Big-data storage", "HDFS /earthscape", "RAW, INTERIM, PROCESSED layers and provenance manifests"],
        ["Batch processing", "Hadoop MapReduce (Streaming) on YARN", "hourly to daily, monthly and yearly aggregation"],
        ["Analytics", "pandas, scikit-learn, SciPy", "trends, anomalies, correlation, forecasts"],
        ["Collection", "Python urllib polling thread", "Open-Meteo weather/air quality, OpenAQ"],
        ["Monitoring and backup", "manage.py health-check / backup, Windows Task Scheduler", "15-minute health check, daily backup with retention"],
    ], [3.4, 6.2, 7.0])

    d.h("6. Data sources and provenance")
    d.table(["Source", "Class", "Notes"], [
        ["NASA POWER hourly", "Gridded reanalysis (MERRA-2) + satellite-derived radiation (CERES); not station data", "125 city-year CSV files, 1,095,720 rows; fill value -999 never present in the download"],
        ["Open-Meteo weather", "MODELLED, near-real-time", "temperature, humidity, precipitation, 10 m wind (km/h), pressure (hPa); UTC; 15-minute interval"],
        ["Open-Meteo air quality", "MODELLED (CAMS)", "PM2.5, PM10, CO, NO2, SO2, O3, US AQI; never shown as observed"],
        ["OpenAQ", "OBSERVED (mostly low-cost sensors, sparse)", "PM2.5 only; active monitors within 25 km; location, sensor class, provider and distance kept; API key only from the environment"],
        ["MODIS MOD11A2 v061", "Satellite-derived", "measured plan: 1,515 granules, 11.2 GB (tiles h24v06 2.9 GB, h24v05 4.2 GB, h23v05 4.1 GB); not downloaded"],
    ], [3.6, 5.4, 7.6])
    d.p("Provenance is recorded in append-only manifest files in HDFS (source, request, size, SHA-256, record counts, times, version, YARN application id). Raw data stays provider-native and immutable; corrections happen downstream. "
        "Missing values stay visible (POWER -999 becomes null plus a flag in INTERIM, never an invented value). Modelled and observed environmental data are labelled on every page and never merged.")

    d.h("7. HDFS design")
    d.code("""/earthscape
  _manifest/<source>.jsonl            append-only provenance (last line per object_key wins)
  raw/power_hourly_gridded/city=<c>/year=<y>/*.csv          125 objects
  raw/openmeteo_weather_model/date=YYYY-MM-DD/polls.jsonl   sealed daily
  raw/openmeteo_airquality_model/date=YYYY-MM-DD/polls.jsonl
  raw/openaq_observed/current/date=YYYY-MM-DD/latest.jsonl
  interim/power_hourly/city=<c>/power_hourly_<c>.csv        5 objects
  processed/power_daily | power_monthly | power_yearly/part-00000""")
    d.p("Partitioning is by city and year (historical) and by source and UTC day (live). Ingestion is idempotent: an identical object is adopted, a different one is refused, uploads go to a temporary path, are verified by size and SHA-256 read-back and are then moved into place. "
        "The live RAW rollover (sealing a finished UTC day) was verified with an in-memory HDFS (eight tests: provenance, counts, hash, duplicates, interruption recovery, no overwrite). "
        "No real daily object exists yet: polling began on 2026-10-08, so the first UTC day ends at 05:00 PKT on 2026-10-09.")

    d.h("8. MapReduce processing")
    d.p("Three Hadoop Streaming jobs (Python mapper and reducer, one reducer, exact decimal arithmetic, deterministic output) are submitted to YARN by run_power_job.py. The driver verifies inputs by manifest, checks the YARN final state, verifies the output with an independent recomputation from INTERIM, publishes without overwriting and writes provenance.")
    d.table(["Job", "Input", "Output", "Rows", "YARN application"], [
        ["daily", "5 INTERIM files", "processed/power_daily", "45,655", "application_1791451740685_0004"],
        ["monthly", "daily output", "processed/power_monthly", "1,500", "application_1791451740685_0005"],
        ["yearly extremes", "daily output", "processed/power_yearly", "125", "application_1791451740685_0006"],
    ], [2.8, 3.0, 4.2, 1.8, 4.8])
    d.p("Daily solar irradiance is the mean in the provider unit MJ/hr (no daily total); the daily precipitation total is empty unless all 24 hourly values are valid. The INTERIM data (67 MB) is smaller than one HDFS block, so the jobs demonstrate the architecture; they do not claim big-data scale, and scaling was not demonstrated on this single node.")

    d.h("9. Machine-learning methodology")
    d.p("Code: src/ml/climate_ml.py (version climate_ml_v1, seed 42). The data are POWER aggregates produced by MapReduce; all statements concern this model-based product, not station observations.")
    d.bullets([
        "**Trends:** Theil-Sen slope with a 95% interval and Mann-Kendall test on annual values (n = 25 years).",
        "**Anomalies:** day-of-year median climatology fitted on 2001-2015 only; robust z-score (|z| > 3.5) and an Isolation Forest (200 trees, contamination 0.01) fitted on 2001-2015 and applied to 2016-2025. No labels exist, so precision and recall cannot be computed; a unit test confirms an injected +15 C spike is detected.",
        "**Correlation:** Pearson and Spearman on deseasonalised daily anomalies; association only, not causation.",
        "**Forecast:** monthly values; chronological split (train 2001-2020, test 2021-2025); seasonal climatology and last-year baselines versus a harmonic-trend Ridge model; model chosen by rolling-origin validation inside the training years only; the band is forecast +/- 1.96 x test RMSE (approximate, not a calibrated interval).",
        "**Updating:** manage.py train-ml or the Administrator button stores a new versioned run in MongoDB; training is reproducible.",
    ])
    d.table(["City", "Temperature trend C/decade [95% interval]", "p", "Precipitation mm/decade [95% interval]", "p"], [
        ["Karachi", "0.127 [-0.053, 0.261]", "0.1557", "107.3 [8.6, 230.2]", "0.0223"],
        ["Lahore", "-0.672 [-1.019, -0.264]", "0.0007", "207.5 [124.2, 324.5]", "0.0000"],
        ["Islamabad", "-0.613 [-1.120, -0.269]", "0.0024", "401.2 [239.0, 598.3]", "0.0001"],
        ["Peshawar", "-0.722 [-1.165, -0.428]", "0.0004", "353.8 [234.3, 465.6]", "0.0000"],
        ["Quetta", "-0.028 [-0.398, 0.346]", "0.9447", "124.5 [40.6, 225.9]", "0.0039"],
    ], [2.4, 5.2, 1.6, 5.4, 2.0])
    d.p("Forecast result: the selected model had a test MAE no worse than the seasonal-climatology baseline in 18 of 20 city-variable forecasts; high R2 for temperature mostly reflects the seasonal cycle and precipitation skill is limited. "
        "Reanalysis trends can include artefacts of changing assimilated observations; this is a description of the dataset, not a physical attribution. Full tables: docs/ml-methodology.md.")

    d.h("10. Dashboard, authentication and authorisation")
    d.p("The application has 15 pages. Authentication uses Argon2id password hashes, generic error messages, a 15-minute lockout after five failures, server-side sessions with random tokens in HttpOnly SameSite cookies (only the SHA-256 is stored) and a CSRF token on every state-changing request. "
        "Roles: **Administrator** (everything, plus user management, feedback review, data refresh, manual polling and retraining) and **Analyst** (dashboards, alert rules, acknowledgements, feedback). The last Administrator cannot be demoted or disabled.")
    d.table(["Page", "Content"], [
        ["Overview, Historical Climate, City Comparison", "KPIs, temperature/precipitation/humidity/wind charts, five-city comparison (NASA POWER, labelled as gridded reanalysis)"],
        ["Trends & Extremes, Correlation, Forecast", "Theil-Sen trends, anomaly counts, hot/frost/heavy-rain days from the yearly MapReduce job, correlation matrices, 12-month forecasts with held-out test results"],
        ["Current Data", "freshness and STALE flags, latest modelled weather and air quality, observed OpenAQ monitors, recent history, current-vs-historical comparison"],
        ["Alert Rules, Alert History", "threshold rules, raised alerts, acknowledgement, navigation badge"],
        ["Data Sources", "provider classes, provenance (YARN ids, checksums), Administrator data refresh"],
        ["Feedback, Feedback Review, User Management", "support requests; Administrator review and account administration"],
    ], [5.0, 11.6])
    for f, cap in (("ui_overview.png", "Figure 2. Overview"), ("ui_history.png", "Figure 3. Historical Climate"), ("ui_trends.png", "Figure 4. Trends and Extremes"),
                   ("ui_forecast.png", "Figure 5. Forecast (model-derived, labelled as predicted)"), ("ui_current.png", "Figure 6. Current Data with the current-vs-historical comparison"),
                   ("ui_alerts_history.png", "Figure 7. Alert History"), ("ui_mobile_overview.png", "Figure 8. Overview at 390 px mobile width")):
        d.image(SHOTS / f, cap, 14.0 if f != "ui_mobile_overview.png" else 6.0)

    d.h("11. Near-real-time processing, alerts, comparison and feedback")
    d.p("A polling thread collects Open-Meteo every 15 minutes and OpenAQ hourly, with bounded retries and backoff. Readings are stored in MongoDB without duplicates (unique per source, city, station and time). "
        "Raw provider responses are staged and sealed into HDFS after each UTC day. Failures are recorded and shown, and never stop other sources. This is near-real-time REST polling, not continuous event streaming.")
    d.p("Alerts: rules name a city (or all), a live variable, operator, threshold and severity. After every poll the newest non-stale reading per city (OpenAQ: per monitoring location) is evaluated; one alert is raised per episode, updated while the condition holds, cleared when it ends; users acknowledge alerts. "
        "Stale readings (weather > 60 min, air quality > 3 h, OpenAQ > 24 h) never raise or clear alerts. Email and SMS are not configured.")
    d.p("Current vs historical: Open-Meteo temperature and humidity of the last 24 hours are compared with the NASA POWER climatology for the same UTC hour (mean and standard deviation over +-7 calendar days across 2001-2025, 375 samples per hour for a typical date). "
        "The two sources are different models on different grids, so differences can be model or grid differences; wind (10 m vs 2 m), pressure (hPa vs kPa), precipitation and PM2.5 are deliberately not compared. A stored reference value was recomputed independently from the HDFS INTERIM data and matched exactly.")
    d.p("Feedback: any user submits a category, subject and message; Administrators review and set the status. Messages containing HTML are shown escaped.")

    d.h("12. Security, monitoring and reliability")
    d.bullets([
        "**Active:** Argon2id, server-side sessions, CSRF, role checks on every route, input validation, parameterised/whitelisted MongoDB access, CSP and security headers, safe error pages, secrets only in the environment or a git-ignored .env.",
        "**Monitoring (active):** a Windows scheduled task runs manage.py health-check every 15 minutes and records application, MongoDB, HDFS, live-polling freshness, newest backup, disk space, CPU and memory with UTC timestamps in logs/health.jsonl. It records only; it does not notify or restart anything.",
        "**Backups (active):** a scheduled task runs a daily backup (MongoDB application data, SHA-256 checksums, newest 7 kept). A real backup was restored into an isolated database with identical document counts and users; the scratch database was then dropped. Backups hold password hashes, never plaintext secrets.",
        "**TLS (prepared, not active):** an Apache reverse-proxy configuration was reviewed and corrected (TLS 1.2/1.3, loopback only); Apache is not installed.",
        "**Encryption at rest (not active):** BitLocker status could not be read without Administrator rights; nothing was enabled.",
        "**Uptime:** 99% was not demonstrated; the system runs on one laptop and WSL daemons stop with their session.",
    ])

    d.h("13. Testing results")
    d.p(f"Automated suite: **{tests} tests, all passing** (stdlib unittest, about two minutes, needs only the local MongoDB; app tests use throw-away databases). Areas: NASA POWER ingestion and validation, RAW to INTERIM, MapReduce mappers/reducers, authentication, RBAC, CSRF, MongoDB, feedback, alerts, data cache, dashboards and APIs, provider parsing, live storage, raw sealing, ML, historical reference and comparison, health monitor, backup retention and restore, MODIS reader on a synthetic HDF4 file.")
    d.bullets([
        "**Hadoop on the real cluster:** three YARN jobs SUCCEEDED; each output was recomputed independently from INTERIM (means within the documented 0.0001 bound); repeat runs were skipped.",
        "**Real browser (Chrome):** 52 page views across both roles at 1280 px and 390 px: all charts painted, no failed requests, no horizontal overflow, role restrictions enforced (403 for the Analyst on admin pages), alert acknowledgement, feedback and error states worked. A defect (unknown URLs returned raw JSON) was found and fixed.",
        "**Live providers:** real polls stored weather, air-quality and 89 OpenAQ readings; repeat polls stored no duplicates.",
        "**Operations:** graceful startup/shutdown of the server verified; scheduled health-check and backup runs finished with result 0.",
    ])

    d.h("14. Limitations and requirements not met")
    d.bullets([
        "**MODIS satellite data not ingested** (no NASA Earthdata credentials). The downloader, plan, HDF4 reader (scale factor, fill value, QC flags, geolocation check) and tests exist; the reader has only been run on a synthetic HDF4 file.",
        "**Weather-station records** (named by the SRS) were not part of the approved datasets.",
        "**Streaming:** near-real-time polling replaces event streaming by decision.",
        "**TLS and encryption at rest** are not active; **99% uptime** and **load balancing** are not demonstrated or implemented; scaling is by design only (single node).",
        "**SRS software list:** Apache server, Tableau, Impala server and RStudio are not installed; Anaconda is replaced by JupyterLab; MongoDB Compass is installed but mongosh is not.",
        "POWER is gridded reanalysis; Open-Meteo is modelled; OpenAQ is sparse and mostly low-cost; forecasts are modest beyond the seasonal baseline; anomaly detection has no ground truth.",
        "The demo video has not been recorded.",
    ])

    d.h("15. Installation and execution")
    d.p("Full steps are in the ReadMe. In short (Windows 11, PowerShell, project root): create the venv and install requirements.txt; start MongoDB; start HDFS (and YARN for new jobs) in WSL; run manage.py init-db, create-admin, refresh-data, build-reference and train-ml; start the server with uvicorn on 127.0.0.1:8000.")
    d.code(""".venv\\Scripts\\python.exe -m uvicorn app.main:create_app --factory --app-dir src --host 127.0.0.1 --port 8000
.venv\\Scripts\\python.exe -m unittest discover -s tests""")

    d.h("16. References")
    d.bullets([
        "NASA Langley Research Center POWER Project, power.larc.nasa.gov (MERRA-2 and CERES SYN1deg based parameters).",
        "Open-Meteo weather and air-quality APIs (open-meteo.com); Copernicus Atmosphere Monitoring Service (CAMS).",
        "OpenAQ API v3 (openaq.org).",
        "NASA LP DAAC, MOD11A2 v061 MODIS/Terra Land Surface Temperature/Emissivity 8-Day L3 Global 1 km SIN Grid.",
        "Apache Hadoop 3.4.3 documentation: HDFS, YARN, MapReduce, Hadoop Streaming.",
        "Sen, P. K. (1968). Estimates of the regression coefficient based on Kendall's tau. JASA 63. Mann, H. B. (1945). Nonparametric tests against trend. Econometrica 13. Kendall, M. G. (1975). Rank Correlation Methods.",
        "Iglewicz, B. and Hoaglin, D. (1993). How to Detect and Handle Outliers. ASQC Quality Press.",
        "Liu, F. T., Ting, K. M. and Zhou, Z.-H. (2008). Isolation Forest. IEEE ICDM.",
        "ETCCDI climate change indices (Zhang et al., 2011, WIREs Climate Change 2).",
    ])
    d.save(OUT / "EarthScape_Final_Report.docx")


def readme():
    d = Doc("EarthScape ReadMe")
    d.d.add_heading("ReadMe - EarthScape Climate Agency", 0)
    d.p("Submission package for the Aptech Big Data eProject. This is the ReadMe.doc required by the SRS (see section 7 for the file formats).")
    d.h("1. Contents of the package")
    d.table(["Path", "What it is"], [
        ["submission/EarthScape_Final_Report.docx / .doc", "The final project report (.docx master and legacy .doc copy)"],
        ["submission/ReadMe.docx / ReadMe.doc", "This document (.docx master and legacy .doc copy, as the SRS names ReadMe.doc)"],
        ["src/", "Source code: app (FastAPI), ingestion, processing (batch, MapReduce), ml"],
        ["tests/", "Automated tests (stdlib unittest)"],
        ["notebooks/", "Exploratory analysis notebook (executed)"],
        ["config/", "Cities, Hadoop configuration for this project, Apache TLS configuration (prepared)"],
        ["docs/", "Architecture, data architecture, ML methodology, security, user guide, testing report, limitations, demo checklist, screenshots"],
        ["scripts/", "ZIP packager, report builder, scheduled-task registration"],
        ["requirements.txt, .env.example", "Pinned Python dependencies; environment variable names only (no secrets)"],
    ], [6.0, 10.6])
    d.p("**Intentionally excluded:** .env and any credentials, the virtual environment, data/ (the 1.1 million-row NASA POWER download, caches, staging), HDFS contents, artifacts, logs, database backups, private keys and the original SRS document. The demo video is not included because it has not been recorded.")
    d.h("2. Assumptions")
    d.bullets([
        "Cities and datasets are the approved set; POWER is gridded reanalysis, Open-Meteo is modelled, OpenAQ is observed.",
        "'Real-time streaming' is implemented as near-real-time REST polling; notifications are in-app only.",
        "Roles are Administrator and Analyst. The application is local (127.0.0.1) and not deployed publicly.",
        "Hadoop 3.4.3 runs inside WSL2 Ubuntu 24.04 and is shared with another project; EarthScape uses only the HDFS path /earthscape and its own configuration.",
        "Apache server, Tableau, Impala server and RStudio (SRS software list) are not installed; MODIS needs NASA Earthdata credentials.",
    ])
    d.h("3. Prerequisites (Windows 11, 16 GB RAM)")
    d.bullets(["Python 3.13", "MongoDB Community Server running on localhost:27017", "WSL2 with Ubuntu 24.04 and Hadoop 3.4.3 in /opt/hadoop with Java 17 (docs/environment.md)", "Internet access (provider APIs; Bootstrap and Chart.js load from a CDN)", "Free OpenAQ API key (OPENAQ_API_KEY) for observed PM2.5"])
    d.h("4. Installation")
    d.code(""".venv creation (project root, PowerShell):
python -m venv .venv
.venv\\Scripts\\python.exe -m pip install -r requirements.txt
copy .env.example .env      (then edit .env: OPENAQ_API_KEY=...; never share this file)""")
    d.h("5. Running the system")
    d.numbered([
        "Start MongoDB (Windows service). Keep a WSL window open: wsl -d Ubuntu-24.04 -e sleep infinity",
        "Start HDFS: wsl -d Ubuntu-24.04 -e /opt/hadoop/sbin/start-dfs.sh (YARN steps in docs/environment.md are needed only for new MapReduce runs).",
        "One-time setup: manage.py init-db; manage.py create-admin --username admin (prompts for a password of 12+ characters).",
        "If the HDFS data exist: manage.py refresh-data; manage.py build-reference; manage.py train-ml.",
        "Start the server: .venv\\Scripts\\python.exe -m uvicorn app.main:create_app --factory --app-dir src --host 127.0.0.1 --port 8000",
        "Open http://127.0.0.1:8000 and sign in. Stop the server with Ctrl+C (graceful shutdown was verified).",
    ])
    d.p("All manage.py commands are run as .venv\\Scripts\\python.exe src\\app\\manage.py <command>. Other commands: poll-once, seal-raw, health-check, backup [--hdfs] [--keep N], restore --path <folder> [--target-db name].")
    d.h("6. Rebuilding the big data (the ZIP excludes it)")
    d.numbered([
        "Ingest NASA POWER RAW (125 city-year files, public API, no key): src\\ingestion\\nasa_power_batch.py.",
        "RAW to INTERIM: src\\processing\\batch\\power_interim.py --city <city> for each city.",
        "MapReduce: src\\processing\\mapreduce\\run_power_job.py daily, then monthly, then yearly (HDFS and YARN running).",
        "Details and expected counts (45,655 / 1,500 / 125 rows): docs/data-architecture.md.",
        "MODIS (optional, needs credentials): add EARTHDATA_TOKEN (or username/password) to .env, then modis.py check and the two-granule pilot; do not download the full 11.2 GB without approval.",
    ])
    d.h("7. Converting this ReadMe and the report to .doc")
    d.p("The SRS names the ReadMe as ReadMe.doc. The master files are .docx. Legacy Word 97-2003 copies (ReadMe.doc and EarthScape_Final_Report.doc) were produced from them with the WPS Office installed on the author's machine (File > Save As > Word 97-2003 Document) and re-opened to check paragraphs, tables and figures; "
        "they were not opened in Microsoft Word itself. If your examiner requires Microsoft Word compatibility, open either file in Word and save it again; check pagination and the figures after converting.")
    d.h("8. Tests")
    d.code(".venv\\Scripts\\python.exe -m unittest discover -s tests")
    d.p("Needs only MongoDB (about two minutes). Hadoop-dependent behaviour is tested with fakes; real-cluster checks are described in the report.")
    d.h("9. Test data")
    d.p("Unit tests use small synthetic fixtures defined in the test files (clearly synthetic and never shown in the application). The application itself uses only real provider data. docs/testing-report.md lists the real-data checks.")
    d.h("10. Known gaps")
    d.p("See Section 14 of the report and docs/srs-completion-plan.md: MODIS not ingested, station records not available, TLS/encryption at rest/uptime/load balancing not active, SRS software list (Apache, Tableau, Impala, RStudio) not installed, demo video not recorded.")
    d.save(OUT / "ReadMe.docx")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--tests", type=int, required=True, help="number of passing tests from the last full run")
    final_report(ap.parse_args().tests)
    readme()
    print("written to", OUT)
